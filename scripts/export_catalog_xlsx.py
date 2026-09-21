"""Экспорт каталога в XLSX (группировка по категориям, картинки)."""
import asyncio
import sys
from pathlib import Path
from decimal import Decimal
import openpyxl
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from sqlalchemy import select
from sqlalchemy.orm import selectinload

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from app.database import async_session_maker, init_db
from app.models import Product, ProductCategory, PriceList, ProductPrice
from app.config import settings

async def export():
    await init_db()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Каталог"
    
    # Стили
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    cat_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    
    headers = ["Картинка", "Название", "Артикул", "Цена (с НДС), ₽"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 60
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 20
    
    async with async_session_maker() as s:
        pl = (await s.execute(select(PriceList).where(PriceList.is_default == True))).scalar_one_or_none()
        cats = (await s.execute(select(ProductCategory).where(ProductCategory.parent_id.is_(None)).order_by(ProductCategory.sort_order))).scalars().all()
        
        row = 2
        for cat in cats:
            # Заголовок группы
            c = ws.cell(row=row, column=1, value=cat.name)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            c.fill = cat_fill
            c.font = Font(bold=True)
            row += 1
            
            prods = (await s.execute(select(Product).where(Product.category_id == cat.id, Product.is_active == True).order_by(Product.name))).scalars().all()
            for p in prods:
                ws.row_dimensions[row].height = 60
                
                # Картинка
                if p.image_file and (settings.CATALOG_IMAGES_DIR / p.image_file).exists():
                    img = Image(str(settings.CATALOG_IMAGES_DIR / p.image_file))
                    img.width = 70
                    img.height = 70
                    ws.add_image(img, f'A{row}')
                
                ws.cell(row=row, column=2, value=p.name).alignment = Alignment(vertical="center")
                ws.cell(row=row, column=3, value=p.sku or "")
                
                price = None
                if pl:
                    pp = (await s.execute(select(ProductPrice.price).where(ProductPrice.product_id == p.id, ProductPrice.price_list_id == pl.id))).scalar_one_or_none()
                    price = pp
                ws.cell(row=row, column=4, value=f"{price:,.2f}".replace(",", " ") if price else "по запросу").alignment = Alignment(horizontal="right", vertical="center")
                row += 1
    
    out_path = settings.STORAGE_DIR / "exports" / "catalog_export.xlsx"
    wb.save(out_path)
    print(f"Экспорт готов: {out_path}")

if __name__ == "__main__":
    asyncio.run(export())
